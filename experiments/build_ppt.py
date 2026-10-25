"""Build docs/presentation.pptx (8 slides, same section order as the report)."""
import json
import re

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from common import FIG, ROOT, TAB
from refs import REFS

R = json.loads((TAB / "results.json").read_text())
NAVY, BLUE, GREY, LIGHT = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x00, 0x72, 0xB2), RGBColor(0x55, 0x55, 0x55), RGBColor(0xEE, 0xF3, 0xF8)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
PLACE = re.compile(r"(\[(?:Student name|Roll no|Guide name|Department|Institute|Academic year)\])")


def runs(p, text, size, bold=False, color=None):
    for part in PLACE.split(text):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Calibri"
        if color:
            r.font.color.rgb = color
        if PLACE.fullmatch(part):
            r.font.color.rgb = RGBColor(0, 0, 0)
            rPr = r._r.get_or_add_rPr()
            hl = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}highlight")
            etree.SubElement(hl, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr").set("val", "FFFF00")


def box(slide, x, y, w, h, lines, size=18, bullet=False, color=None, bold=False, align=None, fill=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    if fill:
        tb.fill.solid()
        tb.fill.fore_color.rgb = fill
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        runs(p, ("•  " if bullet else "") + ln, size, bold, color)
        p.space_after = Pt(6)
        if align:
            p.alignment = align
    return tb


def title(slide, text, n):
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    box(slide, 0.5, 0.15, 12, 0.7, [text], 30, color=RGBColor(255, 255, 255), bold=True)
    box(slide, 12.2, 7.0, 1, 0.4, [str(n)], 12, color=GREY, align=PP_ALIGN.RIGHT)


def pic(slide, name, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), **kw)


# 1 title
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
box(s, 0.8, 1.4, 11.7, 1.6, ["Matched-Filter Chirp Ranging", "with CFAR Detection"], 44, color=RGBColor(255, 255, 255), bold=True)
box(s, 0.8, 3.2, 11.7, 0.9, ["Finding echoes in noise by pattern matching, in simulation and live on a laptop"], 22, color=RGBColor(0xCF, 0xDD, 0xEE))
box(s, 0.8, 4.5, 11.7, 2.6, ["Student name: [Student name]     Roll no: [Roll no]", "Guide name: [Guide name]",
                             "Department: [Department]     Institute: [Institute]", "Academic year: [Academic year]"],
    18, color=RGBColor(255, 255, 255))

# 2 introduction
s = prs.slides.add_slide(BLANK)
title(s, "1  Introduction", 2)
box(s, 0.5, 1.3, 6.4, 5.5, [
    "Distance from an echo: R = c·t / 2",
    "Long chirp = lots of energy; matched filter squeezes it into a peak ≈ 1/B wide [7]",
    "Range resolution c/(2B) = 2.14 cm for B = 8 kHz; SNR gain ≈ 26.8 dB",
    "Problem 1: sidelobes (−13 dB) hide weak targets [1], [2]",
    "Problem 2: a fixed threshold gives false alarms → CFAR [3], [4]",
    "One laptop: speaker + mic only cm apart, so targets < 10–20 cm merge with the direct sound",
], 19, bullet=True)
pic(s, "fig0_pulse_compression.png", 7.0, 1.4, w=6.0)
box(s, 7.0, 5.6, 6.0, 1.2, ["Before (top): echoes invisible at −6 dB SNR. After the matched filter (bottom): clear peaks."], 14, color=GREY)

# 3 literature
s = prs.slides.add_slide(BLANK)
title(s, "2  Comparative Study of Literature", 3)
rows = [("Paper", "Year", "Takeaway for this project"),
        ("Azouz et al., enhanced matched filter, Sensors [1] (base)", "2021", "Shape the receive filter to cut LFM sidelobes"),
        ("Tan & Wang, two-step windowing [2] (abstract only)", "2025", "Sidelobes vs SNR trade-off; not implemented"),
        ("SonicPACT ultrasonic ranging [5]", "2020", "Acoustic ranging on commodity devices"),
        ("Pyroomacoustics [11]", "2018", "Room simulation cross-check"),
        ("Weinberg, sliding-window detectors [4]", "2017", "CA / OS CFAR family"),
        ("Richards [7], Richards et al. [8], Skolnik [9]", "2008–14", "Formulas: resolution, gain, detection"),
        ("Rohling, CFAR in clutter [3]", "1983", "CA-CFAR baseline, multiple targets"),
        ("Allen & Berkley, image method [6]", "1979", "Room echo simulator")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.5), Inches(1.3), Inches(12.3), Inches(5.2)).table
for j, w in enumerate((6.0, 1.3, 5.0)):
    tbl.columns[j].width = Inches(w)
for i, r in enumerate(rows):
    for j, t in enumerate(r):
        c = tbl.cell(i, j)
        c.text = ""
        runs(c.text_frame.paragraphs[0], t, 15 if i else 16, bold=(i == 0), color=RGBColor(255, 255, 255) if i == 0 else None)
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else RGBColor(255, 255, 255))
box(s, 0.5, 6.7, 12, 0.5, ["Base paper first, then latest to oldest. Summaries are title/abstract level."], 13, color=GREY)

# 4 objective
s = prs.slides.add_slide(BLANK)
title(s, "3  Objective", 4)
box(s, 0.6, 1.4, 12, 5.4, [
    "Build an LFM chirp generator and FFT matched filter from scratch in NumPy",
    "Convert delay to range (R = c·t/2); report resolution c/(2B) and unambiguous range",
    "Control sidelobes with rectangular / Hamming / Taylor windows; measure PSLR",
    "Implement CA-CFAR and OS-CFAR with threshold set by a false-alarm probability; test two close targets",
    "Measure Pd, false alarms and range error vs SNR by Monte Carlo",
    "Interactive Streamlit app: Simulation and Live one-device mode on the same code",
    "Optional: estimate room length and breadth, with honest limits",
], 22, bullet=True)

# 5 methodology
s = prs.slides.add_slide(BLANK)
title(s, "4  Methodology", 5)
pic(s, "fig_flowchart.png", 0.4, 1.2, w=8.2)
box(s, 8.8, 1.3, 4.2, 5.6, [
    "Chirp: f0 → f0+B in T",
    "Matched filter: IFFT(FFT(rx)·conj(FFT(ref)))",
    "Window on the replica",
    "CFAR: threshold = α · noise level of training cells",
    "CA: mean, α = N(Pfa^(−1/N) − 1) [3]",
    "OS: k-th smallest, k = 0.75 N [3], [4]",
    "Live: direct sound = zero distance → latency cancels",
], 16, bullet=True)

# 6 work done
s = prs.slides.add_slide(BLANK)
title(s, "5  Work Done", 6)
box(s, 0.5, 1.3, 5.6, 5.6, [
    "Tools: Python, NumPy [12], SciPy (Taylor window only) [13], Matplotlib [14], Streamlit [15], sounddevice [16], pytest [17]",
    "src/: chirp, matched filter, windows, CA/OS-CFAR, metrics, audio I/O, room estimation",
    "32 pytest tests, all passing",
    "experiments/: one script per figure and table",
    "Fallback to simulation if no audio device",
], 17, bullet=True)
pic(s, "../screenshots/app_top.png", 6.4, 1.2, h=5.8)

# 7 results
s = prs.slides.add_slide(BLANK)
title(s, "6  Results", 7)
w = R["windows"]; d = R["detection"]; i10 = d["snr_db"].index(-10.0)
pic(s, "fig3_pd_vs_snr.png", 0.3, 1.15, w=4.4)
pic(s, "fig5_masking.png", 5.0, 1.15, w=7.9)
box(s, 0.4, 4.3, 12.5, 3.0, [
    f"PSLR: rectangular {w['rectangular']['pslr_db']:.1f} dB, Hamming {w['hamming']['pslr_db']:.1f}, Taylor {w['taylor']['pslr_db']:.1f}; main lobe {w['rectangular']['width3db_m']*100:.2f} → {w['taylor']['width3db_m']*100:.2f} cm (Taylor), SNR loss {w['taylor']['snr_loss_db']:.2f} dB",
    f"Pd at −10 dB input SNR: CA {d['pd_ca'][i10]:.3f}, OS {d['pd_os'][i10]:.3f}; measured Pfa {d['pfa_measured_ca']:.1e} / {d['pfa_measured_os']:.1e} vs design 1e-4",
    f"Range error std {d['err_std_mm_ca'][i10]:.1f} mm at −10 dB (one cell = 21.4 mm)",
    f"Masking (2 close targets): weak target found CA {R['masking']['ca']['pd_A']:.2f} vs OS {R['masking']['os']['pd_A']:.2f}",
    f"Room size (simulation): {R['room']['sweep_success']}/{R['room']['sweep_n']} random rooms within 5 cm at 50 dB; 0/36 at 20 dB",
    "Live run and real-room check: NOT done (no audio device here) – placeholders",
], 14, bullet=True)

# 8 references
s = prs.slides.add_slide(BLANK)
title(s, "7  References", 8)
sel = [1, 2, 3, 4, 5, 6, 7, 10]
box(s, 0.5, 1.2, 12.4, 5.9, [f"[{n}] {REFS[n]}" for n in sel], 13)
box(s, 0.5, 6.8, 12.4, 0.5, ["Full list of 19 references in the report. [2] cited from its abstract only."], 12, color=GREY)

out = ROOT / "docs" / "presentation.pptx"
prs.save(out)
print("saved", out, "slides:", len(prs.slides))
