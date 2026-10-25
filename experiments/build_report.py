"""Build docs/report.docx from the real results in docs/tables and docs/figures."""
import csv
import json
import re
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from common import FIG, ROOT, TAB
from refs import REFS
from src.cfar import ca_scale, os_scale
from src.pipeline import RadarConfig

R = json.loads((TAB / "results.json").read_text())
CFG = RadarConfig()
OUT = ROOT / "docs" / "report.docx"
PAGES = json.loads((ROOT / "docs" / ".report_pages.json").read_text()) if (ROOT / "docs" / ".report_pages.json").exists() else {}


def csv_rows(name):
    with open(TAB / f"{name}.csv") as f:
        return list(csv.reader(f))


try:
    out = subprocess.run([sys.executable, "-m", "pytest", "-q"], cwd=ROOT, capture_output=True, text=True).stdout
    N_TESTS = re.search(r"(\d+) passed", out).group(1)
    FAILED = "failed" in out
except Exception:
    N_TESTS, FAILED = "?", True

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
for m in ("left_margin", "right_margin"):
    setattr(sec, m, Inches(1.0))
sec.top_margin = sec.bottom_margin = Inches(0.9)
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(11)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12
for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
    h = doc.styles[name]
    h.font.name, h.font.size, h.font.bold = "Calibri", Pt(size), True
    h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    h.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    h.paragraph_format.keep_with_next = True

fig_n = [0]
tab_n = [0]


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def para(text="", bold=False, italic=False, align=None, size=None, style=None, keep=False):
    p = doc.add_paragraph(style=style)
    add_runs(p, text, bold, italic, size)
    if align:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


PLACE = re.compile(r"(\[(?:Student name|Roll no|Guide name|Department|Institute|Academic year)\])")


def add_runs(p, text, bold=False, italic=False, size=None):
    """Text with **bold** segments; the six placeholders are highlighted yellow."""
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        b = bold
        if seg.startswith("**") and seg.endswith("**"):
            seg, b = seg[2:-2], True
        for part in PLACE.split(seg):
            if not part:
                continue
            r = p.add_run(part)
            r.bold, r.italic = b, italic
            if size:
                r.font.size = Pt(size)
            if PLACE.fullmatch(part):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, it)
        p.paragraph_format.space_after = Pt(2)


def numbered(items):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        add_runs(p, it)
        p.paragraph_format.space_after = Pt(2)


def table(header, rows, widths=None, caption=None, size=9, align_num=True):
    if caption:
        tab_n[0] += 1
        para(f"Table {tab_n[0]}: {caption}", bold=True, size=10, keep=True)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(size)
        shade(c, "DCE6F2")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], str(v), size=size)
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
    # repeat header row, keep rows from splitting
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        trPr.append(cs)
    if widths:
        t.autofit = False
        for i, w in enumerate(widths):
            t.columns[i].width = Inches(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    if len(rows) <= 8:                       # keep short tables on one page
        for row in t.rows[:-1]:
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(name, caption, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIG / name), width=Inches(width))
    fig_n[0] += 1
    c = para(f"Figure {fig_n[0]}: {caption}", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)


def h1(text, page_break=True):
    if page_break:
        doc.add_page_break()
    doc.add_heading(text, 1)


# ---------------------------------------------------------------- footer page numbers
def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(9)


add_page_number(sec)

# ================================================================ COVER
for _ in range(4):
    para()
para("Digital Signal Processing Project Report", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Matched-Filter Chirp Ranging with CFAR Detection", bold=True, size=26, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Pattern matching of a known chirp, and constant-false-alarm-rate target detection, "
     "in simulation and live on a laptop", italic=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(3):
    para()
for label, ph in (("Student name", "[Student name]"), ("Roll no", "[Roll no]"), ("Guide name", "[Guide name]"),
                  ("Department", "[Department]"), ("Institute", "[Institute]"), ("Academic year", "[Academic year]")):
    para(f"{label}: {ph}", size=13, align=WD_ALIGN_PARAGRAPH.CENTER)

# ================================================================ INDEX
doc.add_page_break()
doc.add_heading("Index", 1)
INDEX = ["1  Introduction", "2  Comparative Study of Literature", "3  Objective", "4  Methodology",
         "5  Work Done", "6  Results", "7  References"]
for it in INDEX:
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), alignment=2, leader=1)   # right tab, dotted leader
    add_runs(p, f"{it}\t{PAGES.get(it, '')}")
    p.paragraph_format.space_after = Pt(4)

# ================================================================ 1 INTRODUCTION
h1("1  Introduction")
doc.add_heading("1.1  Background", 2)
para("A pulse-echo system measures distance by timing an echo: if a pulse leaves, bounces off a target and returns "
     "after a delay t, the target is at R = c t / 2, where c is the speed of the wave (343 m/s for sound in air). "
     "A short pulse locates a target sharply but carries little energy, so weak echoes vanish in noise; a long pulse "
     "carries energy but blurs the position. Pulse compression resolves this: transmit a long linear frequency "
     "modulated (LFM) pulse, a chirp whose pitch sweeps through a bandwidth B, and correlate the recording with a copy "
     "of the chirp (a matched filter). Each echo then collapses into a sharp peak about 1/B wide, so range resolution "
     "is c/(2B) while the SNR improves by roughly the number of chirp samples [7], [8], [10].")
para("The matched filter has two practical weaknesses. Its output has sidelobes (about -13 dB for an unweighted "
     "LFM pulse) that can hide weak targets next to strong ones; tapering the filter lowers them at the cost of a wider "
     "peak and some SNR [1], [2]. Second, deciding which peaks are targets with a fixed threshold gives either many "
     "false alarms or missed targets when the noise level changes. CFAR detectors estimate the local noise level "
     "from neighbouring cells and scale it, which keeps the false-alarm rate constant [3], [4], [9].")
doc.add_heading("1.2  Problem addressed", 2)
para("This project builds an interactive tool that (i) generates a chirp, (ii) processes a received signal with a "
     "matched filter written from scratch in NumPy, (iii) detects targets with CA-CFAR and OS-CFAR, and (iv) measures "
     "how well this works. The same code runs on simulated data and on live audio, where a laptop plays the chirp through its "
     "speaker and records with its microphone at the same time. An optional stage estimates the length and breadth of a "
     "rectangular room from the wall echoes.")
doc.add_heading("1.3  Key figures", 2)
a = CFG
tbl = [
    ["Speed of sound in air (parameter c)", "343 m/s at about 20 C", "Standard value; a program parameter (changeable in the app)"],
    ["Range resolution c/(2B), B = 8 kHz", f"{a.resolution_m * 100:.2f} cm", "Formula from [7], [8]; computed in this project"],
    ["Maximum unambiguous range c PRI/2, PRI = 60 ms", f"{a.unambiguous_m:.2f} m", "Formula from [7], [8]; computed in this project"],
    ["Time-bandwidth product B T (8 kHz, 10 ms)", f"{8000 * 0.01:.0f}", "Definition in [7]; computed"],
    ["Matched-filter processing gain 10 log10(T fs)", f"{10 * __import__('math').log10(0.01 * 48000):.1f} dB", "Derivation in [7], [10]; computed (T fs = 480 samples)"],
    ["Peak sidelobe, rectangular / Hamming / Taylor-40 dB", f"{R['windows']['rectangular']['pslr_db']:.1f} / {R['windows']['hamming']['pslr_db']:.1f} / {R['windows']['taylor']['pslr_db']:.1f} dB", "Measured in this project (Section 6.2); about -13.3 dB is the sinc value [7]"],
    ["CA-CFAR scale, N = 32 cells, Pfa = 1e-4", f"{ca_scale(1e-4, 32):.2f}", "Rohling's formula N (Pfa^(-1/N) - 1) [3]; computed"],
    ["OS-CFAR scale, N = 32, k = 24, Pfa = 1e-4", f"{os_scale(1e-4, 32, 24):.2f}", "Solved numerically from the OS-CFAR false-alarm equation [3], [4]"],
]
table(["Quantity", "Value", "Source / how obtained"], tbl, widths=[2.5, 1.5, 2.5],
      caption="Key numbers used in this project. No outside usage statistics are claimed; every value is a constant, "
              "a formula result or a measurement from this project, with the source of the formula named.")
doc.add_heading("1.4  What one laptop can and cannot measure", 2)
para("A laptop has a speaker and a microphone only centimetres apart. The direct speaker-to-microphone sound is far "
     "stronger than any echo and arrives within a fraction of a millisecond. Any reflection from the laptop's own screen or "
     "chassis, a few centimetres away, arrives so soon after the direct sound that it lies inside the same range cell "
     f"({a.resolution_m * 100:.1f} cm for B = 8 kHz) and cannot be separated from it. The screen is therefore not a measurable "
     "target and does not break the method, but it hides anything closer than roughly 10-20 cm. The method is "
     "designed for a wall or box 0.5-3 m away. The direct peak is also useful: it is the zero-distance reference "
     "that cancels the unknown audio latency of the computer.")

# ================================================================ 2 LITERATURE
h1("2  Comparative Study of Literature")
para("The base paper is listed first; the remaining works follow from the latest to the oldest. The benefits and future-scope "
     "columns summarise what each source offers this project, as stated by its title and abstract-level description. "
     "No source was re-read in full during this session (no internet access), and [2] was checked at abstract level only. "
     "\"Future scope\" states what remains open relative to this project.", italic=True, size=10)
lit = [
    ["Design and Implementation of an Enhanced Matched Filter for Sidelobe Reduction of Pulsed LFM Radar [1] (BASE PAPER)",
     "2021", "Sensors 21(11), 3835 (open access)",
     "Shows that the receive filter of a pulsed LFM radar can be designed to lower range sidelobes; gives the project its idea (shape the matched filter to trade sidelobes against main lobe and SNR) and its metrics.",
     "Here the idea is reproduced in the simplest form (windowed replica; PSLR, main-lobe width, SNR loss) for acoustic chirps. The paper's own filter design and implementation are not reproduced."],
    ["Two-step windowing for both sidelobe suppression and SNR maximization of radar range profiles [2]",
     "2025", "Signal, Image and Video Processing 19(2) (abstract only)",
     "Indicates that windowing can be organised in two steps to get low sidelobes and keep SNR; relevant to the SNR loss that a single window causes (Section 6.2).",
     "Not implemented; a two-step window is a natural extension of src/windows.py. Only the abstract was checked."],
    ["SonicPACT: An Ultrasonic Ranging Method for the PACT Protocol [5]",
     "2020", "arXiv:2012.04770",
     "An acoustic ranging method for commodity devices (per its title), supporting the idea that speaker and microphone hardware can range.",
     "Two-device and ultrasonic; this project uses one device and audible sound. Live accuracy on real hardware remains to be measured here."],
    ["Pyroomacoustics: a Python package for audio room simulation and array processing algorithms [11]",
     "2018", "IEEE ICASSP",
     "Provides an image-source room simulator; used here as an independent cross-check of the project's own simulator.",
     "Only first arrivals were compared in a unit test; frequency-dependent absorption and furniture are not modelled in this project."],
    ["An Introduction to Radar Sliding Window Detectors [4]",
     "2017", "arXiv:1709.09786",
     "Overview of sliding-window (CFAR) detectors including cell-averaging and ordered-statistic types; guides the detector design and terminology.",
     "Only CA- and OS-CFAR are implemented; other variants (greatest-of, smallest-of, censored) are left out."],
    ["Fundamentals of Radar Signal Processing, 2nd ed. [7]",
     "2014", "McGraw-Hill (book)",
     "Textbook treatment of matched filtering, pulse compression, resolution c/(2B), ambiguity and detection; source of the formulas used.",
     "General reference; no project-specific results to extend."],
    ["Principles of Modern Radar: Basic Principles [8]",
     "2010", "SciTech Publishing (book)",
     "Reference for waveform, matched-filter and detection fundamentals, used to cross-check definitions.",
     "General reference."],
    ["Radar Handbook, 3rd ed. [9]",
     "2008", "McGraw-Hill (handbook)",
     "Broad reference on radar waveforms, sidelobe control and CFAR processing.",
     "General reference."],
    ["Radar CFAR Thresholding in Clutter and Multiple Target Situations [3]",
     "1983", "IEEE Trans. AES, AES-19(4), 608-621",
     "Baseline analysis of CFAR thresholds when clutter and several targets are present; basis for the CA-CFAR formula and the OS-CFAR idea implemented here.",
     "Real acoustic clutter (furniture, multipath) is only partly explored here; a calibrated clutter model is left open."],
    ["Image method for efficiently simulating small-room acoustics [6]",
     "1979", "J. Acoust. Soc. Am. 65(4), 943-950",
     "The image-source method: a rectangular room's reflections are replaced by mirror-image sources, giving exact arrival times; basis of the room simulator.",
     "Ideal frequency-independent walls only; no furniture or diffuse reverberation."],
    ["Matched Filter and Pulse Compression (Scientific Python course chapter) [10]",
     "undated", "TU Berlin course material",
     "Hands-on Python explanation of correlation, matched filtering and pulse compression; used to cross-check the implementation approach.",
     "Teaching material; no detection or CFAR stage."],
]
table(["Title", "Year", "Source", "Benefits", "Future scope"], lit, widths=[1.45, 0.45, 0.95, 1.9, 1.75], size=8,
      caption="Comparative study of the literature (base paper first, then latest to oldest).")

# ================================================================ 3 OBJECTIVE
h1("3  Objective")
bullets([
    "Implement a linear-FM chirp generator and an FFT-based matched filter in NumPy, without calling a library matched filter.",
    "Locate echoes, convert delay to range with R = c t / 2 (c a parameter), and report range resolution c/(2B) and the unambiguous range.",
    "Study sidelobe control with rectangular, Hamming and Taylor windows and measure PSLR and main-lobe width.",
    "Implement CA-CFAR and OS-CFAR with a threshold set from a target false-alarm probability, and compare them, in particular for two close targets.",
    "Measure detection probability, false-alarm rate and range error against SNR by Monte Carlo simulation.",
    "Provide a Streamlit app with a simulation mode and a live one-device mode (laptop speaker and microphone) that share the same DSP code, with a fallback when no audio device exists.",
    "Optional: estimate the length and breadth of a rectangular room, and state its limits honestly.",
])

# ================================================================ 4 METHODOLOGY
h1("4  Methodology")
figure("fig_flowchart.png", "Flow chart of the system: one DSP chain, two data sources, three outputs.")
doc.add_heading("4.1  Signal model and chirp", 2)
para("The transmitted chirp is s(t) = cos(2 pi (f0 t + 0.5 k t^2)) for 0 <= t < T with sweep rate k = B/T, so the frequency moves "
     "from f0 to f0 + B. The received record is r(t) = sum_i a_i s(t - tau_i) + n(t): delayed, scaled copies of the chirp plus "
     "white Gaussian noise. In live mode the first copy is the direct speaker-to-microphone sound. SNR is defined per sample for a "
     "unit-amplitude chirp, SNR = (1/2) / sigma^2; a target of relative strength a has SNR + 20 log10(a).")
doc.add_heading("4.2  Matched filter and range", 2)
para("The matched filter has impulse response h[n] = conj(s[-n]), the time-reversed conjugate of the reference. Because "
     "convolution is multiplication in the frequency domain, the output is y = IFFT( FFT(r) x conj(FFT(s)) ), with zero padding to avoid "
     "circular wrap-around. Keeping only the positive frequencies (and doubling them) makes y the analytic signal, so |y| is "
     "a smooth envelope without carrier ripple. An echo at delay tau appears as a peak at index tau fs. The range of an echo is "
     "R = c (t_echo - t_direct) / 2 in live mode, which cancels the unknown audio latency. Range resolution is c/(2B); the "
     "maximum unambiguous range is c PRI/2 [7].")
doc.add_heading("4.3  Windows", 2)
para("The reference is multiplied by a window before correlating (rectangular, Hamming, or Taylor with 5 near-in lobes at -40 dB [13]). "
     "Windowing lowers the sidelobes but widens the main lobe and loses peak SNR by (sum(w s^2))^2 / (sum(s^2) sum(w^2 s^2)); "
     "this loss is reported per window. This follows the sidelobe-reduction idea of [1], [2] in its simplest software form.")
doc.add_heading("4.4  CFAR detection", 2)
para("Detection works on the power |y|^2, which for noise alone is exponentially distributed. For each cell under test, N = 2 x "
     "(training cells per side) training cells are collected, skipping the guard cells next to it. CA-CFAR sets the threshold to "
     "alpha x mean(training) with alpha = N (Pfa^(-1/N) - 1) [3]. OS-CFAR uses the k-th smallest training value (k = 0.75 N) "
     "and solves prod_{i=0}^{k-1} (N - i)/(N - i + alpha) = Pfa for alpha [3], [4]. Near the ends of the record fewer cells "
     "exist and N and alpha adapt. Matched-filter outputs are correlated between neighbouring samples (the output is oversampled by "
     "fs/B), so training cells are spaced fs/B samples apart and the guard region must cover the main lobe (Section 6.6). Detections "
     "are local maxima above the threshold; a parabola through three samples gives a sub-sample peak position.")
doc.add_heading("4.5  Monte Carlo evaluation", 2)
para("For each SNR a fixed target is simulated with independent noise in many trials. Detection probability Pd is the fraction of "
     "trials with a detection within one range cell (fs/B samples) of the truth; range error is the difference between the "
     "interpolated peak and the true delay in detected trials; false alarms are detections farther than that tolerance from "
     "every target. The measured false-alarm probability per cell is obtained separately from noise-only records.")
doc.add_heading("4.6  Live mode and room stage", 2)
para("Live mode uses sounddevice [16] full-duplex play-and-record: one chirp at t = 0, 0.4 s recorded. The strongest peak is taken "
     "as the direct sound and must exceed the noise by 10 dB or the app warns. The optional room stage simulates a rectangular room with the "
     "image-source method [6], checked against pyroomacoustics [11]; the estimator tries every grouping of detected echoes into "
     "two wall pairs, predicts all image-source echoes up to third order (including the wall-to-wall echo at the sum of a wall pair), and keeps "
     "the grouping that explains the most echoes. Length = larger total, breadth = smaller. Height is not estimated.")

# ================================================================ 5 WORK DONE
h1("5  Work Done")
doc.add_heading("5.1  Tools", 2)
tools = [
    ["Python", "3.13", "Language", "-"],
    ["NumPy", "2.5.3", "Chirp, FFT matched filter, CFAR, Monte Carlo (all core DSP)", "[12]"],
    ["SciPy", "1.18.1", "Taylor window only (scipy.signal.windows.taylor); no matched filter is taken from the library", "[13]"],
    ["Matplotlib", "3.11.2", "All figures in the report and app", "[14]"],
    ["Streamlit", "1.66.1", "Interactive app (app.py)", "[15]"],
    ["sounddevice", "0.5.6", "Live play-and-record (needs PortAudio)", "[16]"],
    ["pytest", "9.1.1", "Automated tests", "[17]"],
    ["pyroomacoustics", "0.10.1", "Cross-check of the room simulator (test only)", "[11]"],
    ["python-docx / python-pptx", "1.2.0 / 1.0.2", "Generate this report and the slides from the real results", "[18], [19]"],
]
table(["Tool", "Version", "Used for", "Ref."], tools, widths=[1.3, 1.0, 3.5, 0.7], caption="Tools used.")
para("All numerical work uses NumPy [12], with SciPy [13] supplying only the Taylor window coefficients. Figures are drawn with Matplotlib [14], "
     "the app is built with Streamlit [15], live audio uses sounddevice [16], tests run under pytest [17], and this report is generated with python-docx [18] "
     "(the slides with python-pptx [19]).")
doc.add_heading("5.2  Implementation steps", 2)
numbered([
    "**Chirp and formulas** (src/chirp.py): LFM generator with a Nyquist check, exact fractional-delay echoes, resolution and unambiguous-range functions.",
    "**Matched filter** (src/matched_filter.py): zero-padded FFT correlation with the time-reversed conjugate reference, analytic output, optional replica window; a slow direct-convolution version is kept only to test it.",
    "**Windows** (src/windows.py) and the window SNR-loss formula.",
    "**CFAR** (src/cfar.py): vectorised CA- and OS-CFAR with per-cell adaptation at record edges, closed-form CA scale and numerically solved OS scale.",
    "**Pipeline** (src/pipeline.py): filter, threshold, local maxima, minimum separation, sub-sample peak, direct-path reference, 'no direct peak' warning.",
    "**Simulation and metrics** (src/simulation.py, src/metrics.py): echoes plus noise; PSLR and main-lobe width; Monte Carlo Pd, range error, false alarms.",
    "**Audio I/O** (src/audio_io.py): device check, full-duplex record, AudioUnavailable fallback.",
    "**Room stage** (src/room.py): Allen-Berkley image sources, noisy room recordings, wall-pair estimator.",
    "**App** (app.py): sliders for all required parameters, plots with plain-language captions, cached computation, quick and full Monte Carlo.",
    f"**Tests** (tests/): {N_TESTS} pytest tests, all required checks plus audio and room tests.",
    "**Experiments** (experiments/): one script per figure/table, writing to docs/figures and docs/tables from fixed random seeds.",
])
doc.add_heading("5.3  The app", 2)
figure("../screenshots/app_top.png", "Screenshot of the Streamlit app (simulation mode, two targets at 1.2 m and 3.0 m, SNR -6 dB, OS-CFAR); upper part of the page.", width=5.6)
para("The app shows the chirp and its rising pitch, the hidden echoes in the raw signal next to the compressed output, the CFAR threshold with "
     "detections labelled in metres, a detected-versus-true table, an interactive 'cell under test' view of guard and training cells, the window comparison, "
     "a detection-probability curve for the current settings, and the optional room panel.")

# ================================================================ 6 RESULTS
h1("6  Results")
para("All numbers come from the scripts in experiments/ with fixed seeds; nothing was edited by hand. Default settings: fs = 48 kHz, f0 = 3 kHz, B = 8 kHz, "
     "T = 10 ms, Taylor window, 16 training and 8 guard cells per side, Pfa = 1e-4.", italic=True, size=10)
doc.add_heading("6.1  Pulse compression", 2)
figure("fig0_pulse_compression.png", "Top: microphone signal at -6 dB SNR; the echoes cannot be seen. Bottom: matched-filter output; the direct sound and two echoes (1.2 m, 3.0 m) are clear peaks.", 5.8)
doc.add_heading("6.2  Windows: sidelobes against main-lobe width (experiments 1 and 2)", 2)
figure("fig1_windows.png", "Matched-filter response to one target for three windows. Left: sidelobes in dB; right: main-lobe zoom.", 6.3)
wr = csv_rows("windows")
table(["Window", "PSLR (dB)", "-3 dB width (cm)", "Null-to-null (cm)", "SNR loss (dB)"], wr[1:], widths=[1.5, 1.1, 1.3, 1.3, 1.2],
      caption="PSLR, main-lobe width and SNR loss per window (B = 8 kHz, one cell = 2.14 cm).")
para(f"The rectangular window gives {R['windows']['rectangular']['pslr_db']:.1f} dB, close to the sinc value of about -13.3 dB [7]. Hamming reaches "
     f"{R['windows']['hamming']['pslr_db']:.1f} dB and Taylor {R['windows']['taylor']['pslr_db']:.1f} dB, a gain of more than 23 dB, while the -3 dB main lobe widens from "
     f"{R['windows']['rectangular']['width3db_m'] * 100:.2f} cm to {R['windows']['hamming']['width3db_m'] * 100:.2f} cm (Hamming) and {R['windows']['taylor']['width3db_m'] * 100:.2f} cm (Taylor), "
     f"and the peak SNR drops by {R['windows']['hamming']['snr_loss_db']:.2f} and {R['windows']['taylor']['snr_loss_db']:.2f} dB. This is the trade-off described in [1], [2].")
doc.add_heading("6.3  Detection probability against SNR (experiment 3)", 2)
figure("fig3_pd_vs_snr.png", f"Pd versus input SNR for CA- and OS-CFAR, single target at 1.5 m, {R['detection']['trials']} trials per point, shaded 95 % binomial band.", 5.4)
rows = csv_rows("pd_range_error_vs_snr")
table(["SNR (dB)", "Pd CA", "Pd OS"], [[r[0], r[1], r[2]] for r in rows[1:]], widths=[1.2, 1.2, 1.2],
      caption="Detection probability against SNR (same data as the figure).")
d = R["detection"]
i10 = d["snr_db"].index(-10.0)
para(f"Pd rises from near 0 at -20 dB to {d['pd_ca'][i10]:.3f} (CA) and {d['pd_os'][i10]:.3f} (OS) at -10 dB and 1.000 at -8 dB. The SNR here is per sample at the input; "
     f"the matched filter adds about 26.8 dB (less the window loss), which is why targets far below the noise are still found. In homogeneous noise OS-CFAR is slightly less sensitive than CA-CFAR "
     f"(by {d['pd_ca'][i10] - d['pd_os'][i10]:.3f} in Pd at -10 dB), the known price of its robustness.")
para(f"Measured false-alarm probability per cell from noise-only records: CA {d['pfa_measured_ca']:.2e}, OS {d['pfa_measured_os']:.2e}, against a design value of {d['pfa_design']:.0e}. "
     "Both are about 1.6 times higher than designed: neighbouring matched-filter samples are not fully independent, so the noise estimate is noisier than the theory assumes. The CFAR property (a "
     "constant rate) holds, but the rate is not exactly the nominal one.")
doc.add_heading("6.4  Range error (experiment 4)", 2)
figure("fig4_range_error.png", "Standard deviation (left) and mean (right) of the range error in runs that detected the target.", 6.0)
table(["SNR (dB)", "Mean error CA (mm)", "Std CA (mm)", "Mean error OS (mm)", "Std OS (mm)"],
      [[r[0], r[3], r[4], r[5], r[6]] for r in rows[1:] if float(r[1]) > 0.05 or float(r[2]) > 0.05], widths=[1.0, 1.4, 1.2, 1.4, 1.2],
      caption="Range error against SNR (only rows with Pd above 0.05; a 'nan' means no detection).")
para("Because the peak position is interpolated between samples, the error is a small fraction of one range cell (21.4 mm): the standard deviation is about 2.2 mm at -10 dB and "
     "falls to about 1.4 mm at -6 dB, and the mean error stays within a fraction of a millimetre. At very low SNR the few detections include noise peaks and the spread is "
     "larger. These are simulated errors for ideal sound-speed knowledge; in real air a 1 percent error in c would add a 1 percent range error.")
doc.add_heading("6.5  Two close targets: CA-CFAR against OS-CFAR (experiment 5)", 2)
figure("fig5_masking.png", "Weak target A at 1.000 m and a 9.5 dB stronger target B 4.3 cm away (two range cells), SNR -8 dB. Left: one record with the two thresholds. Right: Pd over 1000 trials.", 6.3)
mr = csv_rows("masking")
table(["Detector", "Pd, weak target alone", "Pd, weak target with strong neighbour", "Pd, strong target"], mr[1:], widths=[1.3, 1.6, 2.2, 1.3],
      caption="Masking result.")
para(f"Alone, the weak target is found {R['masking']['ca']['pd_A_alone']:.3f} (CA) and {R['masking']['os']['pd_A_alone']:.3f} (OS) of the time. Next to the strong target CA-CFAR "
     f"finds it in {R['masking']['ca']['pd_A']:.3f} of trials, whereas OS-CFAR still finds it in {R['masking']['os']['pd_A']:.3f}. The strong echo falls in CA-CFAR's training cells, inflates the "
     "average and lifts the threshold above the weak target; OS-CFAR's rank-based estimate ignores the few contaminated cells [3], [4]. Both detect the strong target.")
doc.add_heading("6.6  Tests on new settings", 2)
ns = csv_rows("new_settings")
table(["Setting", "Res. (cm)", "PSLR (dB)", "Pd CA", "Pfa CA", "Pd OS", "Pfa OS"],
      [[r[0], r[1], r[2], r[3], r[4], r[6], r[7]] for r in ns[1:]], widths=[2.3, 0.6, 0.7, 0.6, 0.8, 0.6, 0.8], size=8,
      caption=f"Key checks repeated with other settings (single target, SNR {R['new_settings']['snr_db']:.0f} dB, {R['new_settings']['trials']} trials, design Pfa 1e-4 except where noted).")
para("Resolution follows c/(2B) (1.07 cm at 16 kHz, 5.72 cm at 3 kHz); PSLR follows the window; the wide band loses Pd at this SNR because each sample carries less energy per unit "
     "bandwidth (shorter chirp), and the narrow, long chirp gains Pd. Raising Pfa to 1e-2 yields a measured Pfa of about 1e-2, so the design value is followed over two orders of magnitude. Using 8 instead of 16 "
     "training cells per side raises the noise-estimate variance and costs Pd. A 346 m/s sound speed changes the range scale only (resolution 2.16 cm).")
gs = csv_rows("guard_sweep")
table(["Guard cells / side", "Pd CA", "Pfa CA", "Pd OS", "Pfa OS"], gs[1:], widths=[1.4, 1, 1, 1, 1],
      caption="Effect of the number of guard cells (design Pfa 1e-4, SNR -10 dB, 500 trials).")
para("With too few guard cells the target's own main lobe leaks into the nearest training cells. CA-CFAR then raises its own threshold: with 2 guard cells its Pd falls to "
     f"{R['guard_sweep'][0]['ca']['pd']:.3f} and its measured Pfa is {R['guard_sweep'][0]['ca']['pfa']:.1e}. With 8 guard cells the measured Pfa is close to the design value, which is why 8 is the default. "
     "This is an important practical point: the guard region must cover the main lobe, which is wider for windowed filters.")
doc.add_heading("6.7  Live one-device run (experiment 6)", 2)
lr = csv_rows("live_run")
table(lr[0], lr[1:], widths=[0.5, 1.0, 1.2, 0.8, 1.2, 2.0], size=8, caption="Live run against a tape-measured distance: PLACEHOLDER.")
if R["live"].get("performed"):
    para(f"Mean error {R['live']['mean_error_cm']:.1f} cm, standard deviation {R['live']['std_error_cm']:.1f} cm over {len(R['live']['errors_cm'])} runs.")
else:
    para("**No live measurement was made.** The report was produced in a cloud container without an audio device, so the table above is a placeholder. The live chain is covered by a unit test using a fake "
         "full-duplex device (37 ms unknown latency plus a 1.2 m echo; the recovered range is within 5 mm, showing that the latency cancels), but that is a software check, not a real-world result. "
         "To fill the table, run: python experiments/exp6_live.py --distance <tape distance in m> with a hard flat surface in front of the laptop.")
doc.add_heading("6.8  Optional: room length and breadth (experiment 7)", 2)
figure("fig7_room.png", "Top view of a successful and a failed estimate, and the success rate against the SNR of the direct sound (simulation).", 6.4)
rm = R["room"]
table(["Direct-sound SNR (dB)", "Rooms within 5 cm (of 36)", "Share"], [[k, round(v * 36), f"{v * 100:.0f} %"] for k, v in rm["success_vs_snr"].items()],
      widths=[1.8, 2.0, 1.0], caption="Simulated room-size accuracy (6 rooms from 4x3 m to 9x6 m, 6 random laptop positions each).")
r50 = [r for r in csv_rows("room_estimates")[1:] if r[1] == "50"]
table(["True room (m)", "Positions", "Estimates returned", "Both within 5 cm", "Mean |err| length (cm)", "Mean |err| breadth (cm)"], [[r[0], r[2], r[3], r[4], r[5], r[6]] for r in r50],
      widths=[1.5, 0.8, 1.0, 1.0, 1.1, 1.1], size=8, caption="Per room at 50 dB (errors are computed over the successful estimates only).")
para(f"When the estimate succeeds it is accurate to about a centimetre; the median absolute error over {rm['sweep_n']} random rooms at 50 dB was {rm['sweep_median_err_length_cm']:.1f} cm (length) and "
     f"{rm['sweep_median_err_breadth_cm']:.1f} cm (breadth). But only {rm['sweep_success']} of {rm['sweep_n']} random rooms were within 5 cm on both dimensions; success also falls with SNR "
     "and is zero at 20 dB, because the wall-to-wall echoes needed to pair the walls sink into the noise (the wall echoes are 26-45 dB weaker than the direct sound). Failures (middle panel) are typically a missed "
     "far-wall echo, which leads to a wrong pairing, or floor and ceiling echoes confused with walls. The estimate is not trustworthy without checking the 'confident' flag shown in the app.")
para("**Limits:** rectangular empty room; laptop in the open with its speaker-microphone line parallel to a wall; frequency-independent wall reflection (0.8 for walls, 0.3 floor, 0.4 ceiling in the simulation); no furniture "
     "(real furniture adds echoes this method would misread); height is not estimated; one microphone cannot tell left from right, so walls are reported as nearer or farther; wrong results near corners. "
     "**No real room was measured**, so the live-room comparison with a tape is a placeholder.", size=10)
doc.add_heading("6.9  Automated tests", 2)
para(f"The pytest suite has {N_TESTS} tests and {'**some failed**' if FAILED else 'all of them pass'} (about 6 s). They cover: the self-correlation peak at zero lag with height equal to the chirp energy; "
     "FFT filter equal to time-reversed convolution; delays recovered within one sample (and to 3 mm after interpolation); the resolution and unambiguous-range formulas; windows lowering PSLR "
     "by more than 15 dB; CA- and OS-CFAR false-alarm rates on exponential noise (within 15 %) and on real matched-filter noise (within 0.3-1.5 times the design value); detection of a known target and rejection of noise; "
     "OS-CFAR beating CA-CFAR in the masking case; the no-direct-peak warning; the audio fallback and a fake live loop; image-source agreement with pyroomacoustics [11]; and room estimates within 5 cm for three rooms. "
     "The Streamlit app was also run (streamlit run app.py and Streamlit's AppTest) without exceptions, in both modes (live mode shows the fallback message in this container).")
doc.add_heading("6.10  Points to note", 2)
bullets([
    "All results are simulations except the unit tests of the live chain. **No real microphone or room measurement exists yet.**",
    "SNR is per input sample. Matched filtering adds about 27 dB for the default chirp, so -10 dB input SNR is easy; real laptop speakers and microphones have limited bandwidth and non-flat response, so the achievable B is smaller than the simulated 8 kHz.",
    "The measured false-alarm rate was about 1.6 times the design value (Section 6.3) because matched-filter samples are correlated; CFAR keeps the rate constant, not exactly nominal.",
    "OS-CFAR costs a little sensitivity in clean noise but is far more robust to a close strong neighbour; CA-CFAR missed the weak target in every trial of the masking case.",
    "The guard region must cover the main lobe (8 cells here); too small a guard silently disables CA-CFAR detection.",
    "Targets closer than about 10-20 cm (including the laptop's own screen) cannot be separated from the direct sound.",
    "The room stage works only for simple empty rectangular rooms and is far from reliable at low SNR: 79 of 100 random rooms succeeded at 50 dB, and 0 of 36 at 20 dB.",
    "[2] was cited from its abstract only; other sources were cited at title and abstract level (Section 2).",
])

# ================================================================ 7 REFERENCES
h1("7  References")
for n, text in REFS.items():
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.first_line_indent = Inches(-0.4)
    add_runs(p, f"[{n}]\t{text}", size=10)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))

doc.save(OUT)
print("saved", OUT)
# --- citation check: every reference must be cited in the body (not counting the reference list itself)
body = "\n".join(p.text for p in doc.paragraphs)
body += "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
body = body.split("7  References")[0] if "7  References" in body.split("Index")[-1] and False else body
ref_start = [i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "7  References"][-1]
body = "\n".join(p.text for p in doc.paragraphs[:ref_start]) + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
cited = {int(n) for grp in re.findall(r"\[(\d+(?:\s*[,-]\s*\d+)*)\]", body) for n in re.findall(r"\d+", grp)}
missing = [n for n in REFS if n not in cited]
print("references:", len(REFS), "| not cited in body:", missing)
