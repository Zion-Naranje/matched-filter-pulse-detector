"""Build docs/presentation.pptx (8 slides) for the ESC-50 template-matching project."""
import json, re
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt
from common import FIG, ROOT, TAB

R = json.loads((TAB / "results.json").read_text())
REFS = {int(k): v for k, v in json.loads((ROOT / "docs" / "refs.json").read_text()).items()}
NAVY, BLUE, GREY, LIGHT = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x00, 0x72, 0xB2), RGBColor(0x55, 0x55, 0x55), RGBColor(0xEE, 0xF3, 0xF8)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
PLACE = re.compile(r"(\[(?:Student name|Roll no|Guide name|Department|Institute|Academic year|Base paper to be verified by the student)\])")


def runs(p, text, size, bold=False, color=None):
    for part in PLACE.split(text):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Calibri"
        if color:
            r.font.color.rgb = color
        if PLACE.fullmatch(part):
            r.font.color.rgb = RGBColor(0, 0, 0)
            rPr = r._r.get_or_add_rPr()
            hl = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}highlight")
            etree.SubElement(hl, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr").set("val", "FFFF00")


def box(slide, x, y, w, h, lines, size=18, bullet=False, color=None, bold=False, align=None, fill=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    if fill:
        tb.fill.solid()
        tb.fill.fore_color.rgb = fill
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        runs(p, ("•  " if bullet else "") + ln, size, bold, color)
        p.space_after = Pt(6)
        if align:
            p.alignment = align
    return tb


def title(slide, text, n):
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    box(slide, 0.5, 0.15, 12, 0.7, [text], 30, color=RGBColor(255, 255, 255), bold=True)
    box(slide, 12.2, 7.0, 1, 0.4, [str(n)], 12, color=GREY, align=PP_ALIGN.RIGHT)


def pic(slide, name, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), **kw)



CL = {r["class"]: r for r in R["best"] if r["mode"] == "class"}
RP = {r["class"]: r for r in R["best"] if r["mode"] == "repeat"}
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
box(s, 0.8, 1.3, 11.7, 1.8, ["Template Matching of Environmental Sounds", "using Normalised Cross-Correlation"], 40, color=RGBColor(255, 255, 255), bold=True)
box(s, 0.8, 3.3, 11.7, 0.9, ["Finding where a known sound occurs in a longer recording (ESC-50)"], 22, color=RGBColor(0xCF, 0xDD, 0xEE))
box(s, 0.8, 4.5, 11.7, 2.6, ["Student name: [Student name]     Roll no: [Roll no]", "Guide name: [Guide name]",
                             "Department: [Department]     Institute: [Institute]", "Academic year: [Academic year]"], 18, color=RGBColor(255, 255, 255))

s = prs.slides.add_slide(BLANK); title(s, "Index", 2)
items = ["Title", "Index", "Comparative Study of Literature", "Objective", "Methodology (Flow Chart)", "Work Done", "Results", "References"]
for i, it in enumerate(items):
    box(s, 1.2, 1.3 + i * 0.68, 1.0, 0.6, [f"{i + 1:02d}"], 24, color=BLUE, bold=True)
    box(s, 2.4, 1.3 + i * 0.68, 9, 0.6, [it], 24)
box(s, 1.2, 6.8, 11, 0.5, ["Template matching of environmental sounds with normalised cross-correlation (ESC-50)"], 14, color=GREY)

s = prs.slides.add_slide(BLANK); title(s, "Comparative Study of Literature", 3)
rows = [("Paper", "Year", "Takeaway"),
        ("BASE PAPER: Tutul, Buchem, Pinkwart, cross-correlation vs CNN sound detection (arXiv preprint) [1]", "2026", "Normalised cross-correlation + fixed threshold detects a stored sound, no training"),
        ("Nolasco et al., Learning to detect an animal sound from five examples, Ecol. Informatics [2]", "2023", "Few-shot detection; template matching is the traditional baseline"),
        ("Nolasco et al., Few-shot bioacoustic event detection, DCASE 2022 [3]", "2022", "Challenge overview; detecting a sound from few examples is still open"),
        ("Gong, Chung, Glass, AST: Audio Spectrogram Transformer, Interspeech [4]", "2021", "Learned classifier on spectrograms; reports 95.6 % on ESC-50"),
        ("Kong et al., PANNs, IEEE/ACM TASLP [5]", "2020", "Pretrained audio networks transfer to other tasks")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.5), Inches(1.3), Inches(12.3), Inches(4.6)).table
for j, w in enumerate((6.3, 1.3, 4.7)):
    tbl.columns[j].width = Inches(w)
for i, r in enumerate(rows):
    for j, t in enumerate(r):
        c = tbl.cell(i, j); c.text = ""
        runs(c.text_frame.paragraphs[0], t, 14, bold=(i == 0), color=RGBColor(255, 255, 255) if i == 0 else None)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else RGBColor(255, 255, 255))
box(s, 0.5, 6.3, 12.3, 0.9, ["Base paper: arXiv:2607.16299 (July 2026, open-access preprint), checked against its listing; summarised from the abstract only. All papers 2020 or later, latest to oldest."], 13, color=GREY)

s = prs.slides.add_slide(BLANK); title(s, "Objective", 4)
box(s, 8.3, 1.4, 4.7, 5.4, ["Background", "Template matching = slide a known sound along a recording and score every position (NCC [1], [2])",
    "Test data: ESC-50 [6], [7] clips joined with 1 s silences, so true positions are known", "CC BY-NC 3.0 [8]; 5 classes, 16 kHz"], 15, fill=LIGHT)
box(s, 0.6, 1.4, 7.5, 5.4, [
    "Load ESC-50, convert to mono 16 kHz, normalise",
    "Join clips with 1 s silences; record exact spans as ground truth",
    "FFT-based normalised cross-correlation in NumPy; threshold + non-maximum suppression",
    "Detection correct if its centre lies inside a true clip; precision, recall, F1",
    "Streamlit app with a threshold slider and plots",
    "Experiment: P/R/F1 vs threshold, 5 classes, 30 recordings each; best threshold per class",
], 18, bullet=True)

s = prs.slides.add_slide(BLANK); title(s, "Methodology (Flow Chart)", 5)
pic(s, "fig_flow.png", 0.4, 1.3, w=12.5)
box(s, 0.6, 5.2, 12, 1.8, ["score[k] = sum x[k+n]·t[n] / ( sqrt(sum x[k+n]²) · sqrt(sum t²) )   (numerator by FFT, window energy by running sum)",
                           "Two test conditions: other clips of the class  |  the template itself (gain 0.5–1.5, 20 dB noise)"], 17, bullet=True)

s = prs.slides.add_slide(BLANK); title(s, "Work Done", 6)
box(s, 0.5, 1.3, 5.8, 5.6, [
    "NumPy [10] for all matching; SciPy [11] only for WAV reading and resampling",
    "Data from your ESC-50 zip, a folder, or single-file GitHub download",
    "Streamlit app [13]; Matplotlib [12]",
    "pytest [14]: 6 tests, all pass",
    "Chance baseline added to judge the results honestly",
], 18, bullet=True)
pic(s, "../screenshots/app_repeat.png", 6.5, 1.3, h=5.4)

s = prs.slides.add_slide(BLANK); title(s, "Results", 7)
pic(s, "fig_prf_class.png", 0.3, 1.1, w=12.7)
box(s, 0.4, 4.2, 12.5, 3.1, [
    "Same sound repeated (gain 0.5–1.5, 20 dB noise): F1 = " + ", ".join(f"{c} {float(RP[c]['F1']):.2f}" for c in RP) + " (no false detections)",
    "Other clips of the class (figure): best F1 " + ", ".join(f"{c} {float(CL[c]['F1']):.2f} (chance {float(CL[c]['chance_F1']):.2f})" for c in CL),
    "Only dog and door_wood_knock clearly beat chance; rain and siren are at chance level",
    "Best thresholds are very low (0.02–0.14): raw-waveform NCC does not generalise across different recordings",
], 15, bullet=True)

s = prs.slides.add_slide(BLANK); title(s, "References", 8)
box(s, 0.5, 1.2, 12.4, 5.5, [f"[{n}] {REFS[n]}" for n in (1, 2, 3, 4, 5, 6, 7, 8)], 12)
box(s, 0.5, 6.8, 12.4, 0.5, ["Full list of 17 references in the report. [1] is the base paper; [6] is cited only as the dataset source."], 12, color=GREY)
prs.save(ROOT / "docs" / "presentation.pptx")
print("slides", len(prs.slides))
